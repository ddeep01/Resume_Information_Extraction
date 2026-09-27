import logging
from pathlib import Path
from docx import Document

logger = logging.getLogger("DOCXExtractor")

class DOCXExtractor:
    MIN_TEXT_LENGTH = 50

    @classmethod
    def extract_docx(cls, docx_path: Path) -> tuple[str, str]:
        try:
            doc = Document(str(docx_path))
            text_parts = []

            # 1. Paragraphs
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    text_parts.append(txt)

            # 2. Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            full_text = "\n".join(text_parts).strip()
            if len(full_text) >= cls.MIN_TEXT_LENGTH:
                return full_text, "python-docx"
            return full_text, "python-docx_short" if full_text else "failed"

        except Exception as e:
            logger.error(f"DOCX extraction error for {docx_path}: {e}")
            return "", "failed"
